import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = r"C:\Users\atpad\.gemini\antigravity-ide\scratch\CampusNotify-Java-Mini-Project"

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=True, bottom=True, left=True, right=True, color="000000", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name, active in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if active:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False, color_rgb=(0, 0, 0)):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT, size=12, bold=False, italic=False, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic)
    return p

def add_college_header_table(doc):
    t_hdr = doc.add_table(rows=1, cols=3)
    t_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_mid, c_right = t_hdr.rows[0].cells
    c_left.width = Inches(1.2)
    c_mid.width = Inches(4.3)
    c_right.width = Inches(1.2)
    
    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_path = os.path.join(BASE_DIR, "assets", "college_logo.png")
    if os.path.exists(logo_path):
        p_l.add_run().add_picture(logo_path, width=Inches(1.05))
        
    p_m = c_mid.paragraphs[0]
    p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_m.paragraph_format.line_spacing = 1.15
    p_m.paragraph_format.space_after = Pt(2)
    
    r1 = p_m.add_run("VASANTDADA PATIL PRATISHTHAN'S\nCOLLEGE OF ENGINEERING AND VISUAL ARTS\n")
    set_font(r1, size=11, bold=True)
    r2 = p_m.add_run("(An Autonomous Institute)\nNAAC 'A' Grade Accredited | ISO 9001:2015 Certified\n")
    set_font(r2, size=9, bold=False, italic=True)
    r3 = p_m.add_run("DEPARTMENT OF INFORMATION TECHNOLOGY\n")
    set_font(r3, size=10.5, bold=True)
    r4 = p_m.add_run("NBA Accredited (Dated 01/07/2024 to 30/06/2027)")
    set_font(r4, size=8.5, italic=True)
    
    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    it_logo_path = os.path.join(BASE_DIR, "assets", "it_dept_logo.png")
    if os.path.exists(it_logo_path):
        p_r.add_run().add_picture(it_logo_path, width=Inches(1.05))
        
    for c in [c_left, c_mid, c_right]:
        set_cell_borders(c, top=False, bottom=True, left=False, right=False, color="000000", sz="8")
        set_cell_margins(c, 20, 40, 20, 20)

from generate_all_documents_and_pdfs import build_problem_statement_docx, build_logbook_docx, build_project_report_docx

if __name__ == "__main__":
    ps_path = os.path.join(BASE_DIR, "CampusNotify_Problem_Statement_Synopsis.docx")
    lb_path = os.path.join(BASE_DIR, "CampusNotify_Mini_Project_Logbook.docx")
    rp_path = os.path.join(BASE_DIR, "CampusNotify_Project_Report.docx")
    build_problem_statement_docx(ps_path)
    build_logbook_docx(lb_path)
    build_project_report_docx(rp_path)
    print("Generated all 3 docx documents via generate_docx_documents.py")
